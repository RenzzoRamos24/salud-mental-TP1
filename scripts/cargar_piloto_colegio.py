"""
Carga los 3 packs del piloto en el colegio (A/B/C) al sistema Sami.

Para cada alumno ficticio en los Excel:
  1. Crea User (rol estudiante) + Consent (versión actual).
  2. Crea AplicacionCuestionario apuntando a la plantilla del pack.
  3. Inserta RespuestaAplicacion por cada ítem contestado.
  4. Marca la aplicación como `completado` y llama al EvaluatorService,
     que a su vez dispara BETO (Pack C, frases) y SVM (Pack B, DASS-21).

Idempotencia: si el email del alumno ya existe → reutiliza el usuario y salta
si ya tiene una aplicación para esa plantilla. Podés re-correr sin romper nada.

Uso:
    venv/bin/python -m scripts.cargar_piloto_colegio                # todos
    venv/bin/python -m scripts.cargar_piloto_colegio --pack B       # solo B
    venv/bin/python -m scripts.cargar_piloto_colegio --sin-evaluar  # saltear
                                                                    # eval BETO/SVM
"""
from __future__ import annotations

import argparse
import json
import sys
import uuid
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.security import hash_password
from app.config import settings
from app.database import SessionLocal
from app.models.bank import (
    AplicacionCuestionario,
    BankInstrumento,
    PlantillaBloque,
    PlantillaCuestionario,
    RespuestaAplicacion,
)
from app.models.consent import Consent
from app.models.user import User


# ── Constantes ─────────────────────────────────────────────────────────
DOCS = Path(__file__).resolve().parent.parent / "docs" / "piloto_colegio"

PACK_A_XLSX = DOCS / "respuestas_pack_A_con_nombres.xlsx"
PACK_B_XLSX = DOCS / "respuestas_pack_B_con_nombres.xlsx"
PACK_C_XLSX = DOCS / "respuestas_pack_C_con_nombres.xlsx"

PSICOLOGA_EMAIL = "psicologa@demo.pe"
PASSWORD_ALUMNOS = "Piloto12345"  # temporal — al levantar prod los alumnos cambian

PLANTILLA_A_NOMBRE = "Piloto Colegio · Pack A (PHQ-A + GAD-7)"
PLANTILLA_B_NOMBRE = "Piloto Colegio · Pack B (DASS-21)"
PLANTILLA_C_NOMBRE = "Piloto Colegio · Pack C (Frases SSCT)"

# Mapa Likert texto → num para Google Forms (Pack A)
LIKERT_03 = {
    "Nunca": 0,
    "Algunos días": 1,
    "Más de la mitad": 2,
    "Casi todos los días": 3,
}
LIKERT_DASS = {
    "No me aplicó": 0,
    "Un poco": 1,
    "Bastante": 2,
    "Mucho": 3,
}

# Mapa: frase-piloto (1-20 en el papel) → número en el banco (1-40).
# El banco tiene 40 frases con 5 por área; el piloto usó 20 (2-3 por área).
FRASE_PILOTO_TO_BANK = {
    1: 1,    # En mi casa yo…              → familia
    2: 5,    # Mis hermanos…               → familia
    3: 6,    # Yo soy…                     → autoconcepto
    4: 7,    # Lo que más me gusta de mí…  → autoconcepto
    5: 8,    # Cuando me miro al espejo…   → autoconcepto
    6: 11,   # El colegio para mí…         → escuela
    7: 15,   # Estudiar es…                → escuela
    8: 16,   # Mis amigos…                 → pares
    9: 17,   # Cuando estoy con otros…     → pares
    10: 21,  # Cuando me enojo…            → emociones
    11: 22,  # Cuando estoy triste yo…     → emociones
    12: 25,  # Cuando algo me preocupa…    → emociones
    13: 26,  # Lo que más me da miedo es… → miedos
    14: 28,  # Cuando estoy solo/a…        → miedos
    15: 31,  # Dentro de 5 años…           → futuro
    16: 33,  # Mi mayor sueño es…          → futuro
    17: 34,  # El futuro me…               → futuro
    18: 36,  # Si pudiera cambiar algo…    → identidad
    19: 39,  # Cuando pienso en quién soy… → identidad
    20: 40,  # Lo que define quién soy…    → identidad
}

FRASES_AREAS_TODAS = (
    "familia,autoconcepto,escuela,pares,emociones,miedos,futuro,identidad"
)


# ── Helpers ────────────────────────────────────────────────────────────
def _get_psicologa(db: Session) -> User:
    psi = db.query(User).filter_by(email=PSICOLOGA_EMAIL).first()
    if psi is None:
        sys.exit(
            f"❌ No existe la psicóloga '{PSICOLOGA_EMAIL}'. "
            "Creala primero (registro en la app o script)."
        )
    return psi


def _get_instrumento(db: Session, codigo: str) -> BankInstrumento:
    inst = db.query(BankInstrumento).filter_by(codigo=codigo).first()
    if inst is None:
        sys.exit(f"❌ Falta instrumento '{codigo}' en el banco. Corré los seeds.")
    return inst


def _get_or_create_plantilla(
    db: Session, nombre: str, psi_id: str, descripcion: str,
) -> PlantillaCuestionario:
    pl = db.query(PlantillaCuestionario).filter_by(nombre=nombre).first()
    if pl:
        return pl
    pl = PlantillaCuestionario(
        psicologo_id=psi_id, nombre=nombre, descripcion=descripcion, activa=1,
    )
    db.add(pl)
    db.flush()
    return pl


def _asegurar_plantillas(db: Session, psi_id: str) -> dict:
    """Crea (si no existen) las 3 plantillas del piloto con sus bloques."""
    phqa = _get_instrumento(db, "PHQ-A")
    gad7 = _get_instrumento(db, "GAD-7")
    dass21 = _get_instrumento(db, "DASS-21")

    pl_a = _get_or_create_plantilla(
        db, PLANTILLA_A_NOMBRE, psi_id,
        "Piloto en el colegio — PHQ-A (depresión) + GAD-7 (ansiedad).",
    )
    if not pl_a.bloques:
        db.add(PlantillaBloque(
            plantilla_id=pl_a.id, orden=1, tipo="instrumento",
            instrumento_id=phqa.id,
        ))
        db.add(PlantillaBloque(
            plantilla_id=pl_a.id, orden=2, tipo="instrumento",
            instrumento_id=gad7.id,
        ))

    pl_b = _get_or_create_plantilla(
        db, PLANTILLA_B_NOMBRE, psi_id,
        "Piloto en el colegio — DASS-21. El SVM de Sami da segunda opinión.",
    )
    if not pl_b.bloques:
        db.add(PlantillaBloque(
            plantilla_id=pl_b.id, orden=1, tipo="instrumento",
            instrumento_id=dass21.id,
        ))

    pl_c = _get_or_create_plantilla(
        db, PLANTILLA_C_NOMBRE, psi_id,
        "Piloto en el colegio — 20 frases incompletas SSCT (8 áreas). "
        "BETO clasifica cada respuesta en categorías emocionales.",
    )
    if not pl_c.bloques:
        db.add(PlantillaBloque(
            plantilla_id=pl_c.id, orden=1, tipo="frases",
            frases_areas=FRASES_AREAS_TODAS,
        ))

    db.commit()
    return {"A": pl_a, "B": pl_b, "C": pl_c}


def _get_or_create_alumno(
    db: Session, nombres: str, apellidos: str, email: str, edad: int | None,
    genero: str | None, psi_id: str,
) -> tuple[User, bool]:
    """Devuelve (user, created)."""
    u = db.query(User).filter_by(email=email).first()
    if u:
        return u, False
    u = User(
        id=str(uuid.uuid4()),
        email=email,
        hashed_password=hash_password(PASSWORD_ALUMNOS),
        nombre=nombres,
        apellido=apellidos,
        role="estudiante",
        activo=True,
        psicologo_id=psi_id,
        grado="Secundaria",
        estado_caso="activo",
    )
    db.add(u)
    db.flush()

    # Consentimiento aceptado (para que aparezca en el dashboard clínico)
    db.add(Consent(
        user_id=u.id,
        version=settings.CONSENT_VERSION_ACTUAL,
        aceptado_en=datetime.utcnow() - timedelta(days=1),
        ip_address="10.0.0.1",
    ))
    return u, True


def _existe_aplicacion(db: Session, alumno_id: str, plantilla_id: int) -> bool:
    return db.query(AplicacionCuestionario).filter_by(
        estudiante_id=alumno_id, plantilla_id=plantilla_id,
    ).first() is not None


def _cerrar_y_evaluar(
    db: Session, aplicacion: AplicacionCuestionario, evaluar: bool,
) -> None:
    """Marca completado y (opcionalmente) llama al evaluator."""
    aplicacion.estado = "completado"
    aplicacion.iniciada_at = aplicacion.iniciada_at or (
        datetime.utcnow() - timedelta(minutes=15)
    )
    aplicacion.completada_at = datetime.utcnow()
    db.flush()

    if not evaluar:
        db.commit()
        return

    from app.services.evaluator_service import EvaluatorService
    try:
        resultado = EvaluatorService.evaluar(db, aplicacion)
        aplicacion.resultado_json = json.dumps(resultado, ensure_ascii=False)
        aplicacion.riesgo_global = resultado.get("riesgo_global")
        aplicacion.crisis_activada = bool(resultado.get("crisis_activada"))
    except Exception as e:
        print(f"    ⚠ evaluator falló: {e}")
    db.commit()


# ── Loaders por pack ───────────────────────────────────────────────────
def _leer_xlsx(path: Path) -> list[dict]:
    """Devuelve filas como dicts con las claves = header de la primera fila."""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    filas = []
    for r in range(2, ws.max_row + 1):
        fila = {}
        for c, h in enumerate(headers, start=1):
            fila[h] = ws.cell(row=r, column=c).value
        # descarta filas vacías
        if any(v is not None for v in fila.values()):
            filas.append(fila)
    return filas


def _cargar_pack_A(
    db: Session, plantilla: PlantillaCuestionario, psi_id: str, evaluar: bool,
) -> dict:
    """Pack A = 16 ítems (9 PHQ-A + 7 GAD-7), valores Likert texto."""
    filas = _leer_xlsx(PACK_A_XLSX)
    stats = {"creados": 0, "reusados": 0, "saltados": 0, "cargados": 0, "errores": 0}
    print(f"→ Pack A: {len(filas)} alumnos")

    for i, fila in enumerate(filas, start=1):
        try:
            nombres = str(fila.get("Nombres") or "").strip()
            apellidos = str(fila.get("Apellidos") or "").strip()
            email = str(fila.get("Email") or "").strip()
            edad = fila.get("Edad")
            if not email or not nombres:
                stats["saltados"] += 1
                continue

            alumno, creado = _get_or_create_alumno(
                db, nombres, apellidos, email, edad, None, psi_id,
            )
            stats["creados" if creado else "reusados"] += 1

            if _existe_aplicacion(db, alumno.id, plantilla.id):
                stats["saltados"] += 1
                continue

            app = AplicacionCuestionario(
                plantilla_id=plantilla.id,
                estudiante_id=alumno.id,
                psicologo_id=psi_id,
                estado="pendiente",
                asignada_at=datetime.utcnow() - timedelta(days=1),
            )
            db.add(app)
            db.flush()

            # PHQ-A: ítems 1-9  |  columnas empiezan con "Cómo te has sentido ["
            phqa_cols = [
                h for h in fila.keys()
                if isinstance(h, str) and h.startswith("Cómo te has sentido [")
            ]
            for numero, col in enumerate(phqa_cols[:9], start=1):
                val = fila.get(col)
                if val is None:
                    continue
                num = LIKERT_03.get(str(val).strip())
                if num is None:
                    continue
                db.add(RespuestaAplicacion(
                    aplicacion_id=app.id,
                    origen=f"INSTR:PHQ-A:{numero}",
                    valor_num=num,
                ))

            # GAD-7: ítems 1-7  |  columnas empiezan con "Preocupación y tensión ["
            gad_cols = [
                h for h in fila.keys()
                if isinstance(h, str) and h.startswith("Preocupación y tensión [")
            ]
            for numero, col in enumerate(gad_cols[:7], start=1):
                val = fila.get(col)
                if val is None:
                    continue
                num = LIKERT_03.get(str(val).strip())
                if num is None:
                    continue
                db.add(RespuestaAplicacion(
                    aplicacion_id=app.id,
                    origen=f"INSTR:GAD-7:{numero}",
                    valor_num=num,
                ))

            _cerrar_y_evaluar(db, app, evaluar)
            stats["cargados"] += 1
            print(f"  [{i:>3}/{len(filas)}] {nombres} {apellidos}"
                  f" — riesgo={app.riesgo_global or '—'}")
        except Exception as e:
            db.rollback()
            stats["errores"] += 1
            print(f"    ⚠ fila {i}: {e}")

    return stats


def _cargar_pack_B(
    db: Session, plantilla: PlantillaCuestionario, psi_id: str, evaluar: bool,
) -> dict:
    """Pack B = 21 ítems DASS-21 con valores 0-3 ya numéricos."""
    filas = _leer_xlsx(PACK_B_XLSX)
    stats = {"creados": 0, "reusados": 0, "saltados": 0, "cargados": 0, "errores": 0}
    print(f"→ Pack B: {len(filas)} alumnos")

    for i, fila in enumerate(filas, start=1):
        try:
            nombres = str(fila.get("nombres") or "").strip()
            apellidos = str(fila.get("apellidos") or "").strip()
            email = str(fila.get("email") or "").strip()
            edad = fila.get("edad")
            if not email or not nombres:
                stats["saltados"] += 1
                continue

            alumno, creado = _get_or_create_alumno(
                db, nombres, apellidos, email, edad, None, psi_id,
            )
            stats["creados" if creado else "reusados"] += 1

            if _existe_aplicacion(db, alumno.id, plantilla.id):
                stats["saltados"] += 1
                continue

            app = AplicacionCuestionario(
                plantilla_id=plantilla.id,
                estudiante_id=alumno.id,
                psicologo_id=psi_id,
                estado="pendiente",
                asignada_at=datetime.utcnow() - timedelta(days=1),
            )
            db.add(app)
            db.flush()

            for numero in range(1, 22):
                val = fila.get(f"dass_{numero:02d}")
                if val is None:
                    continue
                try:
                    n = int(val)
                except (TypeError, ValueError):
                    continue
                if n < 0 or n > 3:
                    continue  # -1 = ilegible, no cargamos
                db.add(RespuestaAplicacion(
                    aplicacion_id=app.id,
                    origen=f"INSTR:DASS-21:{numero}",
                    valor_num=n,
                ))

            _cerrar_y_evaluar(db, app, evaluar)
            svm_info = ""
            if evaluar and app.resultado_json:
                r = json.loads(app.resultado_json)
                svm = r.get("svm_segunda_opinion")
                if svm:
                    svm_info = f" | SVM={svm['clase']} (p={svm['probabilidad']:.2f})"
            stats["cargados"] += 1
            print(f"  [{i:>3}/{len(filas)}] {nombres} {apellidos}"
                  f" — riesgo={app.riesgo_global or '—'}{svm_info}")
        except Exception as e:
            db.rollback()
            stats["errores"] += 1
            print(f"    ⚠ fila {i}: {e}")

    return stats


def _cargar_pack_C(
    db: Session, plantilla: PlantillaCuestionario, psi_id: str, evaluar: bool,
) -> dict:
    """Pack C = 20 frases incompletas SSCT en texto libre."""
    filas = _leer_xlsx(PACK_C_XLSX)
    stats = {"creados": 0, "reusados": 0, "saltados": 0, "cargados": 0, "errores": 0}
    print(f"→ Pack C: {len(filas)} alumnos (con BETO — puede tardar ~10 min)")

    for i, fila in enumerate(filas, start=1):
        try:
            nombres = str(fila.get("nombres") or "").strip()
            apellidos = str(fila.get("apellidos") or "").strip()
            email = str(fila.get("email") or "").strip()
            edad = fila.get("edad")
            if not email or not nombres:
                stats["saltados"] += 1
                continue

            alumno, creado = _get_or_create_alumno(
                db, nombres, apellidos, email, edad, None, psi_id,
            )
            stats["creados" if creado else "reusados"] += 1

            if _existe_aplicacion(db, alumno.id, plantilla.id):
                stats["saltados"] += 1
                continue

            app = AplicacionCuestionario(
                plantilla_id=plantilla.id,
                estudiante_id=alumno.id,
                psicologo_id=psi_id,
                estado="pendiente",
                asignada_at=datetime.utcnow() - timedelta(days=1),
            )
            db.add(app)
            db.flush()

            for numero_piloto in range(1, 21):
                texto = fila.get(f"frase_{numero_piloto:02d}")
                if texto is None or str(texto).strip() == "":
                    continue
                bank_numero = FRASE_PILOTO_TO_BANK[numero_piloto]
                db.add(RespuestaAplicacion(
                    aplicacion_id=app.id,
                    origen=f"FRASE:{bank_numero}",
                    valor_texto=str(texto).strip(),
                ))

            _cerrar_y_evaluar(db, app, evaluar)
            stats["cargados"] += 1
            print(f"  [{i:>3}/{len(filas)}] {nombres} {apellidos}"
                  f" — riesgo={app.riesgo_global or '—'}"
                  f" — crisis={'sí' if app.crisis_activada else 'no'}")
        except Exception as e:
            db.rollback()
            stats["errores"] += 1
            print(f"    ⚠ fila {i}: {e}")

    return stats


# ── Entry point ────────────────────────────────────────────────────────
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--pack", choices=["A", "B", "C", "todos"], default="todos",
        help="Cargar solo un pack o todos (default).",
    )
    parser.add_argument(
        "--sin-evaluar", action="store_true",
        help="Cargar respuestas pero NO ejecutar el evaluator (BETO+SVM). "
             "Útil para pruebas rápidas o si querés correr el evaluator después.",
    )
    args = parser.parse_args()
    evaluar = not args.sin_evaluar

    db = SessionLocal()
    try:
        psi = _get_psicologa(db)
        print(f"Psicóloga responsable: {psi.email} (id={psi.id})")

        plantillas = _asegurar_plantillas(db, psi.id)
        print(f"Plantillas: A=id{plantillas['A'].id}  "
              f"B=id{plantillas['B'].id}  C=id{plantillas['C'].id}")
        print()

        summary = {}
        if args.pack in ("A", "todos"):
            summary["A"] = _cargar_pack_A(db, plantillas["A"], psi.id, evaluar)
            print()
        if args.pack in ("B", "todos"):
            summary["B"] = _cargar_pack_B(db, plantillas["B"], psi.id, evaluar)
            print()
        if args.pack in ("C", "todos"):
            summary["C"] = _cargar_pack_C(db, plantillas["C"], psi.id, evaluar)
            print()

        print("=" * 60)
        print("RESUMEN")
        print("=" * 60)
        for pack, s in summary.items():
            print(f"  Pack {pack}:  "
                  f"cargados={s['cargados']}  "
                  f"nuevos={s['creados']}  "
                  f"reusados={s['reusados']}  "
                  f"saltados={s['saltados']}  "
                  f"errores={s['errores']}")
    finally:
        db.close()


if __name__ == "__main__":
    main()
