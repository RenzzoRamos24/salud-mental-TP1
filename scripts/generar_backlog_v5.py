# -*- coding: utf-8 -*-
"""
Genera Product_Backlog_Sprints_v5.xlsx a partir de la v4.

Cambios respecto de la v4 (detalle en scripts/criterios_v5.py):
  1. Seis HU nuevas, HU-76 … HU-81, en un Sprint 12 (octubre 2026).
  2. HU-09 corregida: 4 categorías con softmax, no 8 multietiqueta.
  3. HU-48 y HU-49 marcadas como «Descartada», con la razón en Notas.
  4. Épica nueva EPICA0010 (validación clínica ciega) y EPICA0007 renombrada.
  5. Hoja nueva «Cambios v5» con la trazabilidad de cada cambio.

Reutiliza la lógica de generar_backlog_v4 (no la copia): le cambia el origen,
el destino y la fuente de criterios. La v4 y su generador quedan intactos.

Uso:  venv/bin/python -m scripts.generar_backlog_v5
"""
import shutil
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import scripts.generar_backlog_v4 as base
from scripts import criterios_v5 as v5

ORIGEN = Path("Product_Backlog_Sprints_v4.xlsx")
DESTINO = Path("Product_Backlog_Sprints_v5.xlsx")
GRIS = "7F7F7F"


def preparar_origen(tmp: Path) -> None:
    """Copia la v4 y aplica los ajustes que el generador base lee de la hoja."""
    shutil.copy(ORIGEN, tmp)
    wb = openpyxl.load_workbook(tmp)

    ws = wb["Product Backlog"]
    for r in range(4, ws.max_row + 1):
        aj = v5.AJUSTES_V5.get(str(ws.cell(r, 2).value or "").strip())
        if aj:
            if "historia" in aj:
                ws.cell(r, 4).value = aj["historia"]
            if "notas" in aj:
                ws.cell(r, 9).value = aj["notas"]

    # El generador base toma el objetivo de cada sprint de las filas existentes.
    sp = wb["Sprint Planning"]
    fila = sp.max_row + 1
    sp.cell(fila, 1, v5.SPRINT_12)
    sp.cell(fila, 2, v5.OBJETIVO_SPRINT_12)
    rs = wb["Resumen Sprints"]
    fila = rs.max_row + 1
    rs.cell(fila, 1, v5.SPRINT_12)
    rs.cell(fila, 2, v5.FOCO_SPRINT_12)

    ep = wb["Épicas"]
    filas_ep = {
        str(ep.cell(r, 1).value).strip(): r
        for r in range(4, ep.max_row + 1)
        if str(ep.cell(r, 1).value or "").startswith("EPICA")
    }
    for epica, nombre in v5.EPICAS_V5.items():
        if epica in filas_ep:
            ep.cell(filas_ep[epica], 2).value = nombre
        else:
            fila = ep.max_row + 1
            ep.cell(fila, 1, epica)
            ep.cell(fila, 2, nombre)
    wb.save(tmp)


def retocar_destino() -> dict:
    """Marca las descartadas, ajusta el total y agrega la hoja de cambios."""
    wb = openpyxl.load_workbook(DESTINO)
    ws = wb["Product Backlog"]
    total = vigentes = 0
    for r in range(4, ws.max_row + 1):
        hu = str(ws.cell(r, 2).value or "")
        if not hu.startswith("HU-"):
            continue
        total += 1
        if v5.ESTADO.get(hu) == v5.DESCARTADA:
            for c in range(2, 10):
                ws.cell(r, c).font = Font(size=10, color=GRIS, italic=True)
            ws.cell(r, 11).font = Font(size=9.5, bold=True, italic=True, color=GRIS)
        else:
            vigentes += 1
    for r in range(4, ws.max_row + 1):
        if str(ws.cell(r, 4).value or "").endswith("historias de usuario"):
            ws.cell(r, 4).value = (f"{total} historias de usuario "
                                   f"({vigentes} vigentes · {total - vigentes} descartadas)")

    if "Cambios v5" in wb.sheetnames:
        del wb["Cambios v5"]
    ch = wb.create_sheet("Cambios v5")
    ch["A1"] = "CAMBIOS DE LA VERSIÓN 5 RESPECTO DE LA VERSIÓN 4 (octubre 2026)"
    ch.merge_cells("A1:D1")
    ch["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ch["A1"].fill = PatternFill("solid", fgColor=base.AZUL_TITULO)
    ch["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ch.row_dimensions[1].height = 26
    for c, h in enumerate(["HU", "Tipo de cambio", "Qué cambió", "Por qué / evidencia"], 1):
        cel = ch.cell(3, c, h)
        cel.font = Font(bold=True, color="FFFFFF", size=10.5)
        cel.fill = PatternFill("solid", fgColor=base.AZUL_HEADER)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in zip("ABCD", (9, 16, 62, 70)):
        ch.column_dimensions[col].width = w

    filas = [(h["id"], "Nueva", h["historia"], h["notas"]) for h in v5.NUEVAS_HU_V5]
    filas += [
        ("HU-09", "Corregida", v5.AJUSTES_V5["HU-09"]["historia"],
         "El texto y los criterios decían 8 categorías multietiqueta; el código "
         "usa 4 categorías con softmax y una sola dominante (nlp_service.py, "
         "multi_label=False). También se corrigió el criterio de frase vacía."),
        ("HU-48", "Descartada", "Estado «Descartada».", v5.AJUSTES_V5["HU-48"]["notas"]),
        ("HU-49", "Descartada", "Estado «Descartada».", v5.AJUSTES_V5["HU-49"]["notas"]),
    ]
    for i, vals in enumerate(filas, 4):
        for c, v in enumerate(vals, 1):
            cel = ch.cell(i, c, v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(size=10, color="222222")
            cel.border = base.BORDE
        ch.cell(i, 1).font = Font(size=10, bold=True, color="1F4E79")
    ch.freeze_panes = "A4"

    wb.save(DESTINO)
    return {"total": total, "vigentes": vigentes}


def main():
    if not ORIGEN.exists():
        raise SystemExit(f"No se encontró {ORIGEN}")
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d) / "origen_v5.xlsx"
        preparar_origen(tmp)
        base.ORIGEN = tmp
        base.DESTINO = DESTINO
        base.CRITERIOS = v5.CRITERIOS
        base.ESTADO = v5.ESTADO
        base.NUEVAS_HU = v5.NUEVAS_HU_V5
        base.main()
    r = retocar_destino()
    print(f"{DESTINO}: {r['total']} HU ({r['vigentes']} vigentes, "
          f"{r['total'] - r['vigentes']} descartadas)")


if __name__ == "__main__":
    main()
