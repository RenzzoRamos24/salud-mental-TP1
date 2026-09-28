"""
Convierte docs/manual/MANUAL_USUARIO.md a un .docx con formato de tesis.

Uso:
    venv/bin/python scripts/generar_manual_usuario.py

Salida:
    docs/manual/Manual_Usuario_Sami.docx
"""
import re
import pathlib

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = pathlib.Path(__file__).resolve().parent.parent
ORIGEN = RAIZ / "docs" / "manual" / "MANUAL_USUARIO.md"
DESTINO = RAIZ / "docs" / "manual" / "Manual_Usuario_Sami.docx"

VERDE = RGBColor(0x2F, 0x7D, 0x6E)
TINTA = RGBColor(0x1A, 0x1A, 0x1A)
GRIS = RGBColor(0x6B, 0x72, 0x80)


# ── utilidades de formato ───────────────────────────────────────────────────

def _sombrear(celda, hex_color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _limpiar_inline(texto):
    """Devuelve una lista de (fragmento, negrita, codigo)."""
    texto = texto.replace("&amp;", "&")
    # quita enlaces markdown dejando el rotulo
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    partes = []
    patron = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
    for trozo in patron.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**"):
            partes.append((trozo[2:-2], True, False))
        elif trozo.startswith("`") and trozo.endswith("`"):
            partes.append((trozo[1:-1], False, True))
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            partes.append((trozo[1:-1], True, False))
        else:
            partes.append((trozo, False, False))
    return partes


def _escribir(par, texto, negrita_base=False, color=None, tam=None):
    for frag, neg, cod in _limpiar_inline(texto):
        run = par.add_run(frag)
        run.bold = neg or negrita_base
        if cod:
            run.font.name = "Consolas"
            run.font.size = Pt((tam or 10.5) - 1)
            run.font.color.rgb = VERDE
        else:
            if tam:
                run.font.size = Pt(tam)
            if color is not None:
                run.font.color.rgb = color
    return par


# ── construccion del documento ──────────────────────────────────────────────

def construir():
    md = ORIGEN.read_text(encoding="utf-8")
    doc = Document()

    # estilos base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for sec in doc.sections:
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)

    lineas = md.split("\n")
    i = 0
    n = len(lineas)
    primera_h1 = True
    en_codigo = False
    buffer_codigo = []

    while i < n:
        ln = lineas[i]
        s = ln.strip()

        # ── bloques de codigo ──
        if s.startswith("```"):
            if en_codigo:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(10)
                run = p.add_run("\n".join(buffer_codigo))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x44, 0x44)
                buffer_codigo = []
                en_codigo = False
            else:
                en_codigo = True
            i += 1
            continue
        if en_codigo:
            buffer_codigo.append(ln)
            i += 1
            continue

        # ── separador ──
        if s == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D8E0DE")
            bd.append(bottom)
            pPr.append(bd)
            i += 1
            continue

        # ── tablas ──
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:\-|]+\|$", lineas[i + 1].strip()):
            filas = []
            while i < n and lineas[i].strip().startswith("|"):
                fila = lineas[i].strip()
                if not re.match(r"^\|[\s:\-|]+\|$", fila):
                    celdas = [c.strip() for c in fila.strip("|").split("|")]
                    filas.append(celdas)
                i += 1
            if filas:
                ncols = max(len(f) for f in filas)
                tabla = doc.add_table(rows=0, cols=ncols)
                tabla.style = "Table Grid"
                tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
                for idx, fila in enumerate(filas):
                    celdas_doc = tabla.add_row().cells
                    for j in range(ncols):
                        txt = fila[j] if j < len(fila) else ""
                        par = celdas_doc[j].paragraphs[0]
                        par.paragraph_format.space_after = Pt(2)
                        par.paragraph_format.space_before = Pt(2)
                        _escribir(par, txt, negrita_base=(idx == 0), tam=9.5)
                        if idx == 0:
                            for r in par.runs:
                                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            _sombrear(celdas_doc[j], "2F7D6E")
                        elif idx % 2 == 0:
                            _sombrear(celdas_doc[j], "F4F8F7")
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # ── encabezados ──
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            nivel = len(m.group(1))
            texto = m.group(2).strip()
            if nivel == 1:
                if not primera_h1:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                primera_h1 = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                _escribir(p, texto, negrita_base=True, color=VERDE, tam=19)
            elif nivel == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                _escribir(p, texto, negrita_base=True, color=TINTA, tam=14)
            elif nivel == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                _escribir(p, texto, negrita_base=True, color=VERDE, tam=11.5)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                _escribir(p, texto, negrita_base=True, color=GRIS, tam=10.5)
            i += 1
            continue

        # ── citas / notas ──
        if s.startswith(">"):
            bloque = []
            while i < n and lineas[i].strip().startswith(">"):
                bloque.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            texto = " ".join(x for x in bloque if x)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "2F7D6E")
            bd.append(left)
            pPr.append(bd)
            _escribir(p, texto, tam=10)
            for r in p.runs:
                if r.font.color.rgb is None:
                    r.font.color.rgb = RGBColor(0x33, 0x44, 0x44)
            continue

        # ── listas numeradas ──
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            _escribir(p, m.group(2))
            i += 1
            continue

        # ── listas con vinetas ──
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            _escribir(p, m.group(1))
            i += 1
            continue

        # ── parrafo normal (une lineas continuas) ──
        if s:
            bloque = [s]
            i += 1
            while i < n:
                sig = lineas[i].strip()
                if (not sig or sig.startswith(("#", ">", "|", "-", "*", "```"))
                        or re.match(r"^\d+\.\s", sig)):
                    break
                bloque.append(sig)
                i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _escribir(p, " ".join(bloque))
            continue

        i += 1

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    return DESTINO


if __name__ == "__main__":
    ruta = construir()
    print(f"Manual generado: {ruta}")
