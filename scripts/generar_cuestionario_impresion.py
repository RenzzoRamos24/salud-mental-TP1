"""
Genera el cuestionario del piloto (PHQ-A + GAD-7 + 10 frases) en .docx,
listo para imprimir como respaldo en papel.

Uso:
    venv/bin/python scripts/generar_cuestionario_impresion.py

Salida:
    docs/piloto_colegio/Cuestionario_Piloto_Sami.docx
"""
import pathlib

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = pathlib.Path(__file__).resolve().parent.parent
DESTINO = RAIZ / "docs" / "piloto_colegio" / "Cuestionario_Piloto_Sami.docx"

VERDE = RGBColor(0x2F, 0x7D, 0x6E)
TINTA = RGBColor(0x1A, 0x1A, 0x1A)
GRIS = RGBColor(0x55, 0x55, 0x55)

OPCIONES_0_3 = ["Nunca (0)", "Algunos días (1)", "Más de la mitad (2)", "Casi todos los días (3)"]

PHQA_ITEMS = [
    "Poco interés o placer en hacer cosas.",
    "Sentirte triste, decaído/a o sin esperanza.",
    "Problemas para dormir, dormir demasiado, o despertarte mucho durante la noche.",
    "Sentirte cansado/a o con poca energía.",
    "Falta de apetito o comer en exceso.",
    "Sentirte mal contigo mismo/a, sentir que eres un/a fracasado/a, o que has fallado a tu familia.",
    "Dificultad para concentrarte en cosas como tareas escolares, leer o ver televisión.",
    "Moverte o hablar tan lento que otras personas lo notan, o estar tan inquieto/a que te mueves mucho más de lo habitual.",
    "Pensar que estarías mejor muerto/a o tener pensamientos de hacerte daño de alguna manera.",
]

GAD7_ITEMS = [
    "Sentirte nervioso/a, ansioso/a o con los nervios de punta.",
    "No poder dejar de preocuparte o no poder controlar tus preocupaciones.",
    "Preocuparte demasiado por diferentes cosas.",
    "Tener dificultad para relajarte.",
    "Estar tan inquieto/a que te resulta difícil quedarte quieto/a.",
    "Irritarte o enojarte con facilidad.",
    "Sentir miedo como si algo terrible fuera a pasar.",
]

FRASES = [
    "En mi casa yo…",
    "Yo soy…",
    "El colegio para mí…",
    "Cuando tengo un examen…",
    "Cuando estoy con otros chicos/as…",
    "Cuando estoy triste yo…",
    "Cuando algo me hace feliz…",
    "Cuando algo me preocupa…",
    "Cuando estoy solo/a…",
    "Dentro de 5 años…",
]


def _sombrear(celda, hex_color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _borde_inferior(parrafo, color="999999", sz="6"):
    p_pr = parrafo._p.get_or_add_pPr()
    borde = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), sz)
    inferior.set(qn("w:space"), "4")
    inferior.set(qn("w:color"), color)
    borde.append(inferior)
    p_pr.append(borde)


def _titulo_seccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"  {texto}")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "2F7D6E")
    p_pr.append(shd)
    return p


def _instruccion(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(texto)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = GRIS


def _tabla_likert(doc, items):
    tabla = doc.add_table(rows=1 + len(items), cols=5)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    anchos = [Cm(9.5), Cm(2.4), Cm(2.4), Cm(2.4), Cm(2.4)]

    encabezado = ["", *OPCIONES_0_3]
    for i, texto in enumerate(encabezado):
        celda = tabla.rows[0].cells[i]
        celda.width = anchos[i]
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        run.bold = True
        run.font.size = Pt(8.5)
        _sombrear(celda, "EEF5F3")

    for fila_i, texto in enumerate(items, start=1):
        celda = tabla.rows[fila_i].cells[0]
        celda.width = anchos[0]
        p = celda.paragraphs[0]
        run = p.add_run(f"{fila_i}. {texto}")
        run.font.size = Pt(10)
        for col in range(1, 5):
            celda = tabla.rows[fila_i].cells[col]
            celda.width = anchos[col]
            p = celda.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("☐")
            run.font.size = Pt(14)
    return tabla


def main():
    doc = Document()

    seccion = doc.sections[0]
    seccion.top_margin = Cm(1.6)
    seccion.bottom_margin = Cm(1.6)
    seccion.left_margin = Cm(1.8)
    seccion.right_margin = Cm(1.8)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    estilo.font.color.rgb = TINTA

    titulo = doc.add_paragraph()
    run = titulo.add_run("Sami — Cuestionario de bienestar")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = VERDE

    subt = doc.add_paragraph()
    run = subt.add_run("Colegio · 4to y 5to de secundaria — Re-encuesta")
    run.font.size = Pt(10)
    run.font.color.rgb = GRIS
    subt.paragraph_format.space_after = Pt(10)

    codigo_p = doc.add_paragraph()
    run = codigo_p.add_run("Código de acceso: SAMI-________-____")
    run.font.size = Pt(11)
    _borde_inferior(codigo_p)
    codigo_p.paragraph_format.space_after = Pt(4)

    intro = doc.add_paragraph()
    run = intro.add_run(
        "Este cuestionario es anónimo para tus compañeros. Solo lo revisa la "
        "psicóloga del colegio. No hay respuestas correctas ni incorrectas — "
        "contesta con lo primero que sientas."
    )
    run.font.size = Pt(9.5)
    run.font.color.rgb = GRIS
    intro.paragraph_format.space_after = Pt(6)

    _titulo_seccion(doc, "Parte 1 — PHQ-A (9 preguntas)")
    _instruccion(doc, "Durante las últimas dos semanas, ¿con qué frecuencia te ha molestado lo siguiente?")
    _tabla_likert(doc, PHQA_ITEMS)

    _titulo_seccion(doc, "Parte 2 — GAD-7 (7 preguntas)")
    _instruccion(doc, "Durante las últimas dos semanas, ¿con qué frecuencia te has sentido afectado/a por los siguientes problemas?")
    _tabla_likert(doc, GAD7_ITEMS)

    doc.add_page_break()

    _titulo_seccion(doc, "Parte 3 — Frases incompletas (10 preguntas)")
    _instruccion(doc, "Completa cada frase con lo primero que se te venga a la mente. No hay respuestas correctas ni incorrectas.")

    for i, frase in enumerate(FRASES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"{i}. {frase}")
        run.bold = True
        run.font.size = Pt(11)

        linea = doc.add_paragraph()
        linea.paragraph_format.space_after = Pt(2)
        run = linea.add_run(" ")
        run.font.size = Pt(11)
        _borde_inferior(linea)

    doc.add_paragraph().paragraph_format.space_before = Pt(12)
    pie = doc.add_paragraph()
    run = pie.add_run(
        "Sami — sistema de evaluación de bienestar estudiantil. PHQ-A (Johnson, Harris, "
        "Spitzer & Williams, 2002) · GAD-7 (Spitzer, Kroenke, Williams & Löwe, 2006) · "
        "Frases incompletas: adaptación del Sacks Sentence Completion Test (Sacks & Levy, "
        "1950), selección de 10 ítems validada con la psicóloga supervisora del proyecto. "
        "Este documento es respaldo en papel — las respuestas reales se registran en el "
        "sistema con el código de acceso de cada alumno/a."
    )
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    print(f"Generado: {DESTINO}")


if __name__ == "__main__":
    main()
