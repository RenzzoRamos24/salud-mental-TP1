# -*- coding: utf-8 -*-
"""
Genera Product_Backlog_Sprints_v4.xlsx a partir de v3.

Cambios respecto de v3:
  1. Product Backlog gana dos columnas: "Criterios de Aceptación (Gherkin)"
     y "Estado (verificado en código)".
  2. Se incorporan 16 HU faltantes (HU-60 … HU-75) halladas por auditoría
     inversa código → backlog.
  3. Nueva hoja "Criterios de Aceptación" con un escenario Gherkin por fila.
  4. Se recalculan Sprint Planning, Resumen Sprints y Épicas.

Uso:  venv/bin/python -m scripts.generar_backlog_v4
"""
from pathlib import Path
import shutil
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from scripts.criterios_aceptacion import CRITERIOS, ESTADO, NUEVAS_HU

ORIGEN = Path("Product_Backlog_Sprints_v3.xlsx")
DESTINO = Path("Product_Backlog_Sprints_v4.xlsx")

AZUL_TITULO = "1F4E79"
AZUL_HEADER = "2E5F8F"
GRIS_SUAVE = "F2F6FA"
BORDE = Border(*[Side(style="thin", color="D9E2EC")] * 4)


def gherkin(hu_id: str) -> str:
    """Bloque Gherkin completo de una HU, listo para una celda."""
    escenarios = CRITERIOS.get(hu_id)
    if not escenarios:
        return ""
    partes = []
    for i, esc in enumerate(escenarios, 1):
        partes.append(f"Escenario {i}: {esc['titulo']}")
        partes.extend(f"  {p}" for p in esc["pasos"])
        if i < len(escenarios):
            partes.append("")
    return "\n".join(partes)


def num_sprint(txt) -> int:
    m = re.search(r"(\d+)", str(txt or ""))
    return int(m.group(1)) if m else 99


def main():
    if not ORIGEN.exists():
        raise SystemExit(f"No se encontró {ORIGEN}")
    shutil.copy(ORIGEN, DESTINO)
    wb = openpyxl.load_workbook(DESTINO)

    # ── 1) Leer las filas existentes del Product Backlog ─────────────
    ws = wb["Product Backlog"]
    filas = []
    for r in range(4, ws.max_row + 1):
        hu = ws.cell(r, 2).value
        if not hu or not str(hu).startswith("HU-"):
            continue
        filas.append({
            "id": str(hu).strip(),
            "rol": ws.cell(r, 3).value,
            "historia": ws.cell(r, 4).value,
            "epica": ws.cell(r, 5).value,
            "sp": ws.cell(r, 6).value,
            "prio": ws.cell(r, 7).value,
            "sprint": ws.cell(r, 8).value,
            "notas": ws.cell(r, 9).value,
        })
    existentes = {f["id"] for f in filas}

    # ── 2) Incorporar las HU faltantes ───────────────────────────────
    agregadas = []
    for hu in NUEVAS_HU:
        if hu["id"] in existentes:
            continue
        filas.append({k: hu[k] for k in
                      ("id", "rol", "historia", "epica", "sp", "prio", "sprint", "notas")})
        agregadas.append(hu["id"])

    # Orden: por sprint, y dentro del sprint conservando el orden de llegada
    for i, f in enumerate(filas):
        f["_orden"] = i
    filas.sort(key=lambda f: (num_sprint(f["sprint"]), f["_orden"]))

    # ── 3) Reescribir la hoja Product Backlog ────────────────────────
    ws.delete_rows(4, ws.max_row)          # deja título (1-2) y cabecera (3)
    ws.cell(3, 10, "Criterios de Aceptación (Gherkin)")
    ws.cell(3, 11, "Estado (verificado en código)")
    for c in (10, 11):
        cel = ws.cell(3, c)
        cel.font = Font(bold=True, color="FFFFFF", size=10.5)
        cel.fill = PatternFill("solid", fgColor=AZUL_HEADER)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions["J"].width = 88
    ws.column_dimensions["K"].width = 18

    sin_criterios = []
    for i, f in enumerate(filas, 1):
        r = 3 + i
        crit = gherkin(f["id"])
        if not crit:
            sin_criterios.append(f["id"])
        valores = [i, f["id"], f["rol"], f["historia"], f["epica"], f["sp"],
                   f["prio"], f["sprint"], f["notas"], crit,
                   ESTADO.get(f["id"], "Sin verificar")]
        for c, v in enumerate(valores, 1):
            cel = ws.cell(r, c, v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(size=10, color="222222")
            cel.border = BORDE
            if i % 2 == 0:
                cel.fill = PatternFill("solid", fgColor=GRIS_SUAVE)
        # Nueva HU en negrita para distinguir lo regularizado
        if f["id"] in agregadas:
            ws.cell(r, 2).font = Font(size=10, bold=True, color="1F4E79")
        est = ESTADO.get(f["id"], "Sin verificar")
        color = {"Implementado": "1E7A4B", "Parcial": "B7791F",
                 "No implementado": "C0392B"}.get(est, "555555")
        ws.cell(r, 11).font = Font(size=9.5, bold=True, color=color)
        ws.cell(r, 11).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.cell(r, 10).font = Font(size=9, color="333333", name="Consolas")

    total_sp = sum(int(f["sp"] or 0) for f in filas)
    r = 4 + len(filas)
    ws.cell(r, 5, "TOTAL").font = Font(bold=True)
    ws.cell(r, 6, total_sp).font = Font(bold=True)
    ws.cell(r, 4, f"{len(filas)} historias de usuario").font = Font(bold=True)

    # ── 4) Hoja "Criterios de Aceptación" (un escenario por fila) ────
    if "Criterios de Aceptación" in wb.sheetnames:
        del wb["Criterios de Aceptación"]
    ca = wb.create_sheet("Criterios de Aceptación", 1)
    ca["A1"] = ("CRITERIOS DE ACEPTACIÓN EN FORMATO GHERKIN — "
                "un escenario por fila (Dado / Cuando / Entonces)")
    ca.merge_cells("A1:F1")
    ca["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ca["A1"].fill = PatternFill("solid", fgColor=AZUL_TITULO)
    ca["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ca.row_dimensions[1].height = 26

    cabecera = ["HU", "Sprint", "Rol", "N.º", "Escenario", "Criterio de aceptación (Gherkin)"]
    for c, h in enumerate(cabecera, 1):
        cel = ca.cell(3, c, h)
        cel.font = Font(bold=True, color="FFFFFF", size=10.5)
        cel.fill = PatternFill("solid", fgColor=AZUL_HEADER)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in zip("ABCDEF", (9, 11, 14, 6, 46, 96)):
        ca.column_dimensions[col].width = w
    ca.freeze_panes = "A4"

    fila = 4
    meta = {f["id"]: f for f in filas}
    for f in filas:
        for i, esc in enumerate(CRITERIOS.get(f["id"], []), 1):
            texto = "\n".join(esc["pasos"])
            vals = [f["id"], f["sprint"], f["rol"], i, esc["titulo"], texto]
            for c, v in enumerate(vals, 1):
                cel = ca.cell(fila, c, v)
                cel.alignment = Alignment(wrap_text=True, vertical="top")
                cel.font = Font(size=10, color="222222")
                cel.border = BORDE
            ca.cell(fila, 6).font = Font(size=9.5, color="333333", name="Consolas")
            ca.cell(fila, 1).font = Font(size=10, bold=True, color="1F4E79")
            ca.cell(fila, 4).alignment = Alignment(horizontal="center", vertical="top")
            if fila % 2 == 0:
                for c in range(1, 7):
                    ca.cell(fila, c).fill = PatternFill("solid", fgColor=GRIS_SUAVE)
            fila += 1
    total_esc = fila - 4

    # ── 5) Recalcular Sprint Planning ────────────────────────────────
    sp_ws = wb["Sprint Planning"]
    objetivos = {}
    for r in range(4, sp_ws.max_row + 1):
        s = sp_ws.cell(r, 1).value
        if s and str(s).lower().startswith("sprint"):
            objetivos[str(s).strip()] = sp_ws.cell(r, 2).value
    sp_ws.delete_rows(4, sp_ws.max_row)

    sprints = sorted({str(f["sprint"]).strip() for f in filas}, key=num_sprint)
    r = 4
    for s in sprints:
        hus = [f for f in filas if str(f["sprint"]).strip() == s]
        vals = [s, objetivos.get(s, ""), ", ".join(h["id"] for h in hus), len(hus),
                sum(int(h["sp"] or 0) for h in hus),
                "\n".join(f"• {h['id']}: {h['historia']}" for h in hus)]
        for c, v in enumerate(vals, 1):
            cel = sp_ws.cell(r, c, v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(size=10, color="222222")
            cel.border = BORDE
        r += 1

    # ── 6) Recalcular Resumen Sprints ────────────────────────────────
    rs = wb["Resumen Sprints"]
    focos = {}
    for r_ in range(4, rs.max_row + 1):
        s = rs.cell(r_, 1).value
        if s and str(s).lower().startswith("sprint"):
            focos[str(s).strip()] = rs.cell(r_, 2).value
    rs.delete_rows(4, rs.max_row)
    r = 4
    for s in sprints:
        hus = [f for f in filas if str(f["sprint"]).strip() == s]
        eps = sorted({str(h["epica"]) for h in hus})
        vals = [s, focos.get(s, ""), len(hus), sum(int(h["sp"] or 0) for h in hus),
                ", ".join(eps)]
        for c, v in enumerate(vals, 1):
            cel = rs.cell(r, c, v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(size=10, color="222222")
            cel.border = BORDE
        r += 1
    epicas_all = sorted({str(f["epica"]) for f in filas})
    for c, v in enumerate(["TOTAL", "", len(filas), total_sp, f"{len(epicas_all)} épicas"], 1):
        cel = rs.cell(r, c, v)
        cel.font = Font(bold=True)
        cel.alignment = Alignment(wrap_text=True, vertical="top")

    # ── 7) Recalcular Épicas ─────────────────────────────────────────
    ep = wb["Épicas"]
    objetivos_ep = {}
    for r_ in range(4, ep.max_row + 1):
        e = ep.cell(r_, 1).value
        if e and str(e).startswith("EPICA"):
            objetivos_ep[str(e).strip()] = ep.cell(r_, 2).value
    ep.delete_rows(4, ep.max_row)
    r = 4
    for e in epicas_all:
        hus = [f for f in filas if str(f["epica"]) == e]
        vals = [e, objetivos_ep.get(e, ""), len(hus), ", ".join(h["id"] for h in hus)]
        for c, v in enumerate(vals, 1):
            cel = ep.cell(r, c, v)
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            cel.font = Font(size=10, color="222222")
            cel.border = BORDE
        r += 1

    wb.save(DESTINO)

    print(f"OK -> {DESTINO}")
    print(f"   HU totales:        {len(filas)}  (v3 tenía {len(existentes)})")
    print(f"   HU agregadas:      {len(agregadas)} -> {', '.join(agregadas)}")
    print(f"   Story points:      {total_sp}")
    print(f"   Escenarios Gherkin:{total_esc}")
    print(f"   Sprints:           {len(sprints)} | Épicas: {len(epicas_all)}")
    if sin_criterios:
        print(f"   ⚠ HU SIN criterios: {sin_criterios}")
    else:
        print("   Todas las HU tienen criterios de aceptación.")


if __name__ == "__main__":
    main()
