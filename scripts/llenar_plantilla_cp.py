# -*- coding: utf-8 -*-
"""
Llena la «Plantilla 02 - Casos de Prueba» oficial del curso.

Estructura exigida por la plantilla:
  · hoja «LISTA CP» — listado con CP, descripción, HU, # escenario y criterio.
  · una hoja por cada caso de prueba, con autor, precondiciones, pasos
    numerados, datos de prueba, resultados esperados, HU relacionada y
    postcondiciones.

Uso:  venv/bin/python -m scripts.llenar_plantilla_cp
"""
from copy import copy
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from scripts.criterios_aceptacion import CRITERIOS
from scripts.generar_hu_cp import leer_backlog, CODIGO
from scripts.llenar_plantilla_hu import id_plantilla, tres_columnas

PLANTILLA = Path("Plantilla 02 - Casos de Prueba.xlsx")
DESTINO = Path(f"{CODIGO}_CP.xlsx")
AUTOR = "Renzo Gutiérrez Ramos"

RE_COMILLAS = re.compile(r'"([^"]{1,60})"')


def datos_de_prueba(texto: str) -> str:
    """Valores literales entrecomillados del escenario, si los hay."""
    vistos, out = set(), []
    for v in RE_COMILLAS.findall(texto):
        if v not in vistos:
            vistos.add(v)
            out.append(v)
    return " · ".join(out)


# Los criterios están redactados en primera persona (es lo natural en una
# historia de usuario). Un caso de prueba se redacta en infinitivo, así que
# los verbos de apertura se convierten al ejecutar.
_INFINITIVO = {
    "Abro": "Abrir", "Intento": "Intentar", "Envío": "Enviar",
    "Solicito": "Solicitar", "Inicio": "Iniciar", "Llamo": "Llamar",
    "Ingreso": "Ingresar", "Creo": "Crear", "Consulto": "Consultar",
    "Pulso": "Pulsar", "Modifico": "Modificar", "Marco": "Marcar",
    "Entro": "Entrar", "Escribo": "Escribir", "Navego": "Navegar",
    "Aplico": "Aplicar", "Ejecuto": "Ejecutar", "Leo": "Leer",
    "Realizo": "Realizar", "Confirmo": "Confirmar", "Completo": "Completar",
    "Indico": "Indicar", "Reviso": "Revisar", "Genero": "Generar",
    "Asigno": "Asignar", "Subo": "Subir", "Descargo": "Descargar",
    "Vinculo": "Vincular", "Guardo": "Guardar", "Vuelvo": "Volver",
    "Selecciono": "Seleccionar", "Elijo": "Elegir",
}
RE_REFLEXIVO = re.compile(r"\b(\w+[aei]r)me\b")
RE_ANONIMO = re.compile(
    r"no autenticad|visitante|sin sesión|sin haber inicia|/register|/login|"
    r"olvidé mi contraseña|token de recuperación", re.IGNORECASE)


def a_infinitivo(paso: str) -> str:
    t = paso.strip()
    primera = t.split(" ", 1)[0]
    if primera in _INFINITIVO:
        resto = t[len(primera):]
        t = _INFINITIVO[primera] + resto
    t = RE_REFLEXIVO.sub(lambda m: m.group(1) + "se", t)
    return t if t.endswith(".") else t + "."


def pasos_del_caso(actor, ctx, evt):
    """Pasos numerados en infinitivo: acceso + acciones del escenario.

    El contexto NO se repite como paso: ya figura en Precondiciones.
    """
    if str(actor).strip().lower() == "sistema":
        pasos = ["Disponer del sistema en ejecución, con el planificador activo."]
    elif RE_ANONIMO.search(ctx):
        pasos = ["Acceder a la aplicación web sin haber iniciado sesión."]
    else:
        pasos = [f"Iniciar sesión en el sistema con un usuario de rol «{actor}»."]
    for linea in evt.split("\n"):
        t = linea.strip()
        if t and t != "N/A":
            pasos.append(a_infinitivo(t))
    return pasos


def main():
    if not PLANTILLA.exists():
        raise SystemExit(f"Falta la plantilla oficial: {PLANTILLA}")

    filas = leer_backlog()

    # ── Construir la lista de casos ──────────────────────────────────
    casos, n = [], 0
    for f in filas:
        for i, esc in enumerate(CRITERIOS[f["id"]], 1):
            n += 1
            ctx, evt, res = tres_columnas(esc["pasos"])
            texto = " ".join(esc["pasos"])
            casos.append({
                "cp": f"CP{n:03d}",
                "desc": esc["titulo"].upper(),
                "hu": id_plantilla(f["id"]),
                "n_esc": i,
                "criterio": esc["titulo"],
                "actor": f["rol"],
                "ctx": ctx, "evt": evt, "res": res,
                "datos": datos_de_prueba(texto),
            })

    wb = openpyxl.load_workbook(PLANTILLA)

    # ── Hoja «LISTA CP» ──────────────────────────────────────────────
    ws = wb["LISTA CP"]
    ws.delete_rows(3, ws.max_row)
    for i, c in enumerate(casos):
        r = 3 + i
        for col, val in ((2, c["cp"]), (3, c["desc"]), (4, c["hu"]),
                         (5, c["n_esc"]), (6, c["criterio"])):
            cel = ws.cell(r, col, val)
            cel.alignment = Alignment(wrap_text=True, vertical="top",
                                      horizontal="center" if col in (2, 4, 5) else "left")
    fin = 3 + len(casos)
    ws.cell(fin + 1, 2, " (*) Proviene de la Plantilla de HISTORIAS DE USUARIO Y "
                        "CRITERIOS DE ACEPTACIÓN")
    ws.cell(fin + 2, 2, "NOTA: Es una hoja por cada CP")
    for rr in (fin + 1, fin + 2):
        ws.cell(rr, 2).font = Font(size=9, italic=True)
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=6)
    for col, w in (("B", 12), ("C", 62), ("D", 14), ("E", 12), ("F", 58)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B3"

    # ── Una hoja por caso ────────────────────────────────────────────
    modelo = wb["CP2"]
    for c in casos:
        hoja = wb.copy_worksheet(modelo)
        hoja.title = c["cp"]

        for rango in list(hoja.merged_cells.ranges):
            hoja.unmerge_cells(str(rango))
        hoja.delete_rows(1, hoja.max_row)

        estilo_tit = copy(modelo["A1"]._style)
        estilo_cab = copy(modelo["A6"]._style)

        hoja.cell(1, 1, f"Caso de Prueba: {c['cp']}: {c['desc']}")
        hoja["A1"]._style = estilo_tit
        hoja.merge_cells("A1:D1")

        hoja.cell(2, 1, "Autor:")
        hoja.cell(2, 2, AUTOR)
        hoja.merge_cells("B2:D2")

        hoja.cell(3, 1, "Precondiciones: " + c["ctx"].replace("\n", "  ·  "))
        hoja.merge_cells("A3:D5")
        hoja["A3"].alignment = Alignment(wrap_text=True, vertical="top")

        for col, txt in ((1, "#:"), (2, "Pasos a seguir:"),
                         (3, "Datos de Prueba"), (4, "Resultados Esperados:")):
            cel = hoja.cell(6, col, txt)
            cel._style = copy(estilo_cab)
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        pasos = pasos_del_caso(c["actor"], c["ctx"], c["evt"])
        r = 7
        for i, paso in enumerate(pasos, 1):
            hoja.cell(r, 1, i).alignment = Alignment(horizontal="center", vertical="top")
            hoja.cell(r, 2, paso).alignment = Alignment(wrap_text=True, vertical="top")
            if i == 1 and c["datos"]:
                hoja.cell(r, 3, c["datos"]).alignment = Alignment(wrap_text=True, vertical="top")
            if i == len(pasos):
                hoja.cell(r, 4, c["res"]).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1

        hoja.cell(r, 1, f"HU relacionada : {c['hu']} — escenario {c['n_esc']} "
                        f"({c['criterio']})")
        hoja.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        hoja.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="center")
        r += 1

        ultima = [x.strip() for x in c["res"].split("\n") if x.strip()]
        post = ultima[-1] if ultima else "El sistema conserva su estado."
        if post.startswith("Y "):
            post = post[2:]
        post = post[:1].upper() + post[1:]
        hoja.cell(r, 1, "Postcondiciones: " + post)
        hoja.merge_cells(start_row=r, start_column=1, end_row=r + 2, end_column=4)
        hoja.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")

        for col, w in (("A", 6), ("B", 66), ("C", 30), ("D", 62)):
            hoja.column_dimensions[col].width = w

    for nombre in ("CP1", "CP2", "CPn"):
        if nombre in wb.sheetnames:
            del wb[nombre]

    wb.save(DESTINO)
    print(f"OK  {DESTINO}")
    print(f"    {len(casos)} casos de prueba · {len(wb.sheetnames)} hojas")
    print(f"    hojas: {wb.sheetnames[:4]} ... {wb.sheetnames[-2:]}")


if __name__ == "__main__":
    main()
