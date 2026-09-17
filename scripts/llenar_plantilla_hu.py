# -*- coding: utf-8 -*-
"""
Llena la «Plantilla 01 - Historias de Usuario y Criterios de Aceptación»
oficial del curso con las 69 HU y los 164 escenarios del proyecto.

Respeta la estructura original: no se agregan ni quitan columnas, se conservan
las hojas «Instructivo» y «Ejemplo», y se mantienen las fórmulas REDACCION
(columna G) y de identificador del criterio (columna J).

Uso:  venv/bin/python -m scripts.llenar_plantilla_hu
"""
from copy import copy
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Alignment

from scripts.criterios_aceptacion import CRITERIOS
from scripts.generar_hu_cp import leer_backlog, EPICAS, num_sprint, CODIGO

PLANTILLA = Path("Plantilla 01 - Historias de Usuario y Criterios de Aceptación_v1.0.xlsx")
DESTINO = Path(f"{CODIGO}_HU.xlsx")

RE_HISTORIA = re.compile(
    r"^Como\s+(?:un[ao]?\s+)?(.+?)\s+quiero\s+(.+?)\s+para\s+(.+)$",
    re.IGNORECASE | re.DOTALL)


def id_plantilla(hu_id: str) -> str:
    """HU-01 -> HU0001, tal como exige la plantilla."""
    return "HU{:04d}".format(int(re.search(r"\d+", hu_id).group()))


def descomponer(historia: str):
    m = RE_HISTORIA.match(str(historia).strip())
    if not m:
        raise ValueError(f"No se pudo descomponer la historia: {historia!r}")
    return tuple(x.strip() for x in m.groups())


def _limpiar(paso: str, *prefijos: str) -> str:
    """Quita la palabra clave Gherkin y deja prosa legible."""
    t = paso.strip()
    for p in prefijos:
        if t.lower().startswith(p.lower()):
            t = t[len(p):].strip()
            break
    return t[:1].upper() + t[1:] if t else t


def tres_columnas(pasos):
    """Gherkin -> (Contexto, Evento, Resultado) según la plantilla."""
    ctx, evt, res = [], [], []
    destino, prefijos = ctx, ("Dado que", "Dado")
    for paso in pasos:
        t = paso.strip()
        if t.startswith("Dado"):
            destino, prefijos = ctx, ("Dado que", "Dado")
        elif t.startswith("Cuando"):
            destino, prefijos = evt, ("Cuando",)
        elif t.startswith("Entonces"):
            destino, prefijos = res, ("Entonces",)
        elif t.startswith("Y "):
            destino.append(t)          # continuación: se conserva tal cual
            continue
        destino.append(_limpiar(t, *prefijos))
    return ("\n".join(ctx) or "N/A",
            "\n".join(evt) or "N/A",
            "\n".join(res) or "N/A")


def clonar_estilo(ws, fila_modelo: int, fila: int, cols: range):
    for c in cols:
        origen = ws.cell(fila_modelo, c)
        destino = ws.cell(fila, c)
        destino._style = copy(origen._style)
        destino.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if not PLANTILLA.exists():
        raise SystemExit(f"Falta la plantilla oficial: {PLANTILLA}")

    filas = leer_backlog()
    wb = openpyxl.load_workbook(PLANTILLA)

    # ── Hoja «HIstoria de Usuario» ───────────────────────────────────
    ws = wb["HIstoria de Usuario"]
    ws.delete_rows(3, ws.max_row)           # conserva las dos filas de cabecera

    r = 3
    for f in filas:
        rol, caracteristica, razon = descomponer(f["historia"])
        hu = id_plantilla(f["id"])
        escenarios = CRITERIOS[f["id"]]
        for n, esc in enumerate(escenarios, 1):
            ctx, evt, res = tres_columnas(esc["pasos"])
            primera = (n == 1)
            ws.cell(r, 1, hu)                                   # A  ID historia
            if primera:
                ws.cell(r, 2, rol)                              # B  Rol
                ws.cell(r, 4, caracteristica)                   # D  Funcionalidad
                ws.cell(r, 6, razon)                            # F  Razón
            ws.cell(r, 3, "Necesito:")                          # C  rótulo
            ws.cell(r, 5, "Con la finalidad de:")               # E  rótulo
            ws.cell(r, 7, f'=CONCATENATE("Como  ",B{r}," quiero ", D{r}," para ",F{r})')
            ws.cell(r, 8, n)                                    # H  # escenario
            ws.cell(r, 9, esc["titulo"])                        # I  título
            ws.cell(r, 10, f'=CONCATENATE(A{r},"-",H{r})')      # J  ID criterio
            ws.cell(r, 11, ctx)                                 # K  Contexto
            ws.cell(r, 12, evt)                                 # L  Evento
            ws.cell(r, 13, res)                                 # M  Resultado
            for c in range(1, 14):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1

    anchos = {"A": 13, "B": 15, "C": 11, "D": 46, "E": 20, "F": 46, "G": 78,
              "H": 9, "I": 40, "J": 14, "K": 52, "L": 46, "M": 56}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"

    # ── Hoja «EPICAS» ────────────────────────────────────────────────
    we = wb["EPICAS"]
    for rango in list(we.merged_cells.ranges):
        we.unmerge_cells(str(rango))
    we.delete_rows(3, we.max_row)

    r = 3
    for eid, objetivo in sorted(EPICAS.items()):
        hus = [x for x in filas if x["epica"] == eid]
        if not hus:
            continue
        inicio = r
        for h in hus:
            we.cell(r, 4, id_plantilla(h["id"]))
            we.cell(r, 4).alignment = Alignment(horizontal="center", vertical="center")
            r += 1
        we.cell(inicio, 2, eid)
        we.cell(inicio, 3, objetivo)
        if r - 1 > inicio:
            we.merge_cells(start_row=inicio, start_column=2, end_row=r - 1, end_column=2)
            we.merge_cells(start_row=inicio, start_column=3, end_row=r - 1, end_column=3)
        for c in (2, 3):
            we.cell(inicio, c).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True)
    we.column_dimensions["B"].width = 20
    we.column_dimensions["C"].width = 56
    we.column_dimensions["D"].width = 26

    wb.save(DESTINO)
    total_esc = sum(len(CRITERIOS[f["id"]]) for f in filas)
    print(f"OK  {DESTINO}")
    print(f"    hojas: {wb.sheetnames}")
    print(f"    {len(filas)} HU · {total_esc} escenarios · {len(EPICAS)} épicas")


if __name__ == "__main__":
    main()
