"""
Genera Product_Backlog_Sprints_v3.xlsx — agrega Sprint 11 (Rol Padre +
firma + descarga en Word) sobre el backlog v2 existente.

Nuevas HUs:
  HU-55 — Padre lista los hijos a su cargo y ve un resumen no clínico.
  HU-56 — Psicóloga sube su firma manuscrita para los informes.
  HU-57 — Psicóloga descarga el reporte clínico individual en Word.
  HU-58 — Padre descarga el informe firmado (PDF y Word).
  HU-59 — Admin vincula/desvincula padres con sus hijos.
"""
from pathlib import Path
import shutil

import openpyxl
from openpyxl.styles import Alignment, Font

ORIGEN = Path("Product_Backlog_Sprints_v2.xlsx")
DESTINO = Path("Product_Backlog_Sprints_v3.xlsx")

NUEVAS_HU = [
    {
        "id": "HU-55", "rol": "Padre",
        "historia": ("Como Padre/Tutor quiero ver el listado de mis hijos vinculados "
                     "con un resumen no clínico (grado, estado, # cuestionarios completados) "
                     "para acompañar su proceso de bienestar"),
        "epica": "EPICA0009", "sp": 5, "prio": "Alta", "sprint": "Sprint 11",
        "notas": ("Endpoint GET /padre/hijos. No se exponen puntajes, banderas ni notas "
                  "internas — solo metadata + nombre de psicóloga. Cumple Ley 29733."),
    },
    {
        "id": "HU-56", "rol": "Psicólogo",
        "historia": ("Como Psicólogo quiero subir mi firma manuscrita (PNG/JPG) "
                     "para que se anexe automáticamente al final de cada informe "
                     "clínico y de los informes que descarguen los padres"),
        "epica": "EPICA0003", "sp": 3, "prio": "Alta", "sprint": "Sprint 11",
        "notas": ("Endpoint POST/DELETE /psychologist/firma + columna users.firma_path. "
                  "Máx 2 MB, PNG o JPG. Se inserta en report_service y word_service."),
    },
    {
        "id": "HU-57", "rol": "Psicólogo",
        "historia": ("Como Psicólogo quiero descargar el reporte clínico individual en "
                     "formato Word (.docx) para poder editarlo o anexarlo a una derivación"),
        "epica": "EPICA0003", "sp": 3, "prio": "Media", "sprint": "Sprint 11",
        "notas": ("Endpoint GET /psychologist/students/{id}/report.docx. Usa python-docx. "
                  "Misma estructura que el PDF de HU-34."),
    },
    {
        "id": "HU-58", "rol": "Padre",
        "historia": ("Como Padre/Tutor quiero descargar un informe oficial de mi hijo "
                     "firmado por la psicóloga (PDF y Word) para archivarlo o llevarlo "
                     "a una atención externa"),
        "epica": "EPICA0009", "sp": 5, "prio": "Alta", "sprint": "Sprint 11",
        "notas": ("Endpoints GET /padre/hijos/{id}/informe.{pdf,docx}. Versión sin datos "
                  "clínicos sensibles. Incluye firma de la psicóloga (HU-56)."),
    },
    {
        "id": "HU-59", "rol": "Administrador",
        "historia": ("Como Administrador quiero vincular y desvincular un padre con sus "
                     "hijos para que el padre pueda acceder al panel correspondiente"),
        "epica": "EPICA0004", "sp": 3, "prio": "Alta", "sprint": "Sprint 11",
        "notas": ("Endpoint POST /admin/padres/{vincular,desvincular}. Un padre puede "
                  "tutelar N estudiantes (1 a muchos vía users.padre_id)."),
    },
]


def main():
    shutil.copy(ORIGEN, DESTINO)
    wb = openpyxl.load_workbook(DESTINO)

    # ── 1) Hoja Product Backlog ──────────────────────────────────────
    ws = wb["Product Backlog"]
    fila = ws.max_row
    # Buscar la fila "TOTAL" y reemplazarla
    fila_total = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 5).value or "").upper().strip() == "TOTAL":
            fila_total = r
            break
    if fila_total:
        ws.delete_rows(fila_total)

    inicio = ws.max_row + 1
    n = ws.max_row - 3 + 1  # número de fila visible (descartando header)
    proximo_num = 48 + 1  # ya había 48 HU
    for i, hu in enumerate(NUEVAS_HU):
        r = ws.max_row + 1
        ws.cell(r, 1, proximo_num + i)
        ws.cell(r, 2, hu["id"])
        ws.cell(r, 3, hu["rol"])
        ws.cell(r, 4, hu["historia"])
        ws.cell(r, 5, hu["epica"])
        ws.cell(r, 6, hu["sp"])
        ws.cell(r, 7, hu["prio"])
        ws.cell(r, 8, hu["sprint"])
        ws.cell(r, 9, hu["notas"])
        for c in range(1, 10):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # Re-agregar TOTAL al final
    r = ws.max_row + 1
    ws.cell(r, 5, "TOTAL")
    sp_total = 238 + sum(h["sp"] for h in NUEVAS_HU)
    ws.cell(r, 6, sp_total)
    ws.cell(r, 5).font = Font(bold=True)
    ws.cell(r, 6).font = Font(bold=True)

    # ── 2) Hoja Sprint Planning ──────────────────────────────────────
    sp = wb["Sprint Planning"]
    r = sp.max_row + 1
    sp.cell(r, 1, "Sprint 11")
    sp.cell(r, 2, ("Rol Padre + Firma de la psicóloga + Descarga en Word\n\n"
                   "Se incorpora el rol Padre/Tutor con acceso restringido a sus hijos. "
                   "La psicóloga firma sus informes una vez. Los reportes clínicos y los "
                   "informes para padre se descargan también en formato Word."))
    sp.cell(r, 3, ", ".join(h["id"] for h in NUEVAS_HU))
    sp.cell(r, 4, len(NUEVAS_HU))
    sp.cell(r, 5, sum(h["sp"] for h in NUEVAS_HU))
    sp.cell(r, 6, "\n".join(f"• {h['id']}: {h['historia']}" for h in NUEVAS_HU))
    for c in range(1, 7):
        sp.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")

    # ── 3) Hoja Resumen Sprints ──────────────────────────────────────
    rs = wb["Resumen Sprints"]
    # Insertar antes de TOTAL
    fila_total_rs = None
    for r in range(1, rs.max_row + 1):
        if str(rs.cell(r, 1).value or "").upper().strip() == "TOTAL":
            fila_total_rs = r
            break

    if fila_total_rs:
        rs.insert_rows(fila_total_rs)
        rs.cell(fila_total_rs, 1, "Sprint 11")
        rs.cell(fila_total_rs, 2, "Rol Padre + firma psicóloga + descarga Word")
        rs.cell(fila_total_rs, 3, len(NUEVAS_HU))
        rs.cell(fila_total_rs, 4, sum(h["sp"] for h in NUEVAS_HU))
        rs.cell(fila_total_rs, 5, "EPICA0003, EPICA0004, EPICA0009")
        # Actualizar fila TOTAL
        fila_total_rs += 1
        rs.cell(fila_total_rs, 3, 48 + len(NUEVAS_HU))
        rs.cell(fila_total_rs, 4, 238 + sum(h["sp"] for h in NUEVAS_HU))
        rs.cell(fila_total_rs, 5, "9 épicas")

    # ── 4) Hoja Épicas — nueva épica EPICA0009 ───────────────────────
    ep = wb["Épicas"]
    fila = ep.max_row + 1
    ep.cell(fila, 1, "EPICA0009")
    ep.cell(fila, 2, "Acceso del padre/tutor al seguimiento del estudiante")
    ep.cell(fila, 3, 2)  # HU-55, HU-58
    ep.cell(fila, 4, "HU-55, HU-58")
    for c in range(1, 5):
        ep.cell(fila, c).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(DESTINO)
    print(f"OK -> {DESTINO}")
    print(f"   Nuevas HU: {len(NUEVAS_HU)}")
    print(f"   Total HU: {48 + len(NUEVAS_HU)}")
    print(f"   Total SP: {238 + sum(h['sp'] for h in NUEVAS_HU)}")


if __name__ == "__main__":
    main()
